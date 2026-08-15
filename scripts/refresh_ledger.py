"""One command for the whole "new Zoho export landed" pipeline: rebuild
Invoice Dump (and Credit Note dump, if given), then bring the new invoice
data into Invoice working (Recurring Revenue / MRR Movement's actual
source) automatically — the three separate manual steps this took the
first time around, chained into one.

Run:
    py -3 scripts/refresh_ledger.py <Invoice export.xlsx> [--credit-notes <CreditNotes export.xlsx|.csv>] [--cutoff YYYY-MM-DD]

--cutoff defaults to the 1st of the invoice export's own earliest invoice
date's month — i.e. "everything this export covers, and only that" — so
it doesn't need to be worked out and typed by hand each time; pass it
explicitly to override.

Nothing here touches DATA_END (frontend/js/core/dates.js) — that one
recomputes itself from the ledger's own data on every page load, on
purpose, so it never needs a code change when new months of data arrive.
"""
import argparse
import subprocess
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MN = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]


def earliest_invdate_month(invoice_xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(invoice_xlsx, read_only=True, data_only=True)
    ws = wb["Invoice"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header.index("Invoice Date")
    earliest = None
    for r in ws.iter_rows(min_row=2, values_only=True):
        v = r[idx]
        if v is None:
            continue
        d = v.date() if hasattr(v, "date") else None
        if d is None:
            continue
        if earliest is None or d < earliest:
            earliest = d
    if earliest is None:
        raise SystemExit("Could not find any Invoice Date in the source file")
    return date(earliest.year, earliest.month, 1)


def run(args_list):
    print("$", " ".join(str(a) for a in args_list))
    subprocess.run([sys.executable] + [str(a) for a in args_list], check=True, cwd=ROOT)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoice_export", help="Path to the Invoice export .xlsx")
    ap.add_argument("--credit-notes", help="Path to the Credit Notes export .xlsx or .csv")
    ap.add_argument("--cutoff", help="YYYY-MM-DD; defaults to the invoice export's own earliest month")
    args = ap.parse_args()

    # Computed once and passed to both merges explicitly, rather than each
    # script re-detecting its own — keeps invoicedump.json and consol.json
    # merged at exactly the same boundary even if their source files don't
    # perfectly agree on where their own data starts.
    cutoff = args.cutoff or earliest_invdate_month(args.invoice_export).isoformat()
    print(f"Using cutoff: {cutoff}")

    run([ROOT / "scripts" / "build_invoice_dump.py", args.invoice_export, "--cutoff", cutoff])
    if args.credit_notes:
        run([ROOT / "scripts" / "build_credit_note_dump.py", args.credit_notes])
    run([ROOT / "scripts" / "merge_consol_from_dump.py", "--cutoff", cutoff])

    print()
    print("Done. Invoice Dump, Credit Note dump (if given), and Invoice working are all refreshed.")
    print("Recurring Revenue / MRR Movement will pick up the new months on next page load --")
    print("DATA_END recomputes itself from the ledger, no code change needed.")


if __name__ == "__main__":
    main()
